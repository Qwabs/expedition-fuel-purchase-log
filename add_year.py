#!/usr/bin/env python3
"""Append one or more years to the fuel purchase master, preserving all data."""

import argparse
import datetime as dt
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

from build_fuel_workbook import OUT, add_years, detect_years


def main():
    parser = argparse.ArgumentParser(
        description="Add a year to Fuel_Purchase_Log_Master.xlsx without touching existing data."
    )
    parser.add_argument("years", nargs="+", type=int, help="Year(s), e.g. 2031 2032")
    parser.add_argument("--file", default=str(OUT), help="Master workbook path")
    args = parser.parse_args()

    path = Path(args.file)
    if not path.exists():
        print(f"Not found: {path}")
        return 1

    years = sorted(set(args.years))
    invalid = [y for y in years if y < 2000 or y > 2100]
    if invalid:
        print(f"Out-of-range year(s): {invalid}")
        return 1

    wb = load_workbook(path)
    existing = detect_years(wb)
    additions = [y for y in years if y not in existing]
    if not additions:
        print(f"Nothing to do: {', '.join(map(str, years))} already present.")
        return 0

    backup_dir = path.parent / "backups"
    backup_dir.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = backup_dir / f"{path.stem}_{stamp}.xlsx"
    shutil.copy2(path, backup)

    try:
        added = add_years(path, additions)
    except Exception:
        shutil.copy2(backup, path)
        raise

    wb_after = load_workbook(path)
    print(f"Backup:  {backup}")
    print(f"Added:   {', '.join(map(str, added))}")
    print(f"Years:   {', '.join(map(str, detect_years(wb_after)))}")
    print(f"Sheets:  {len(wb_after.sheetnames)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
