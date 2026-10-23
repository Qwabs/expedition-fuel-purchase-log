#!/usr/bin/env python3
"""Structural verification for Fuel_Purchase_Log_Master.xlsx (year-log edition)."""

import calendar
import datetime as dt
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

PATH = Path(__file__).with_name("Fuel_Purchase_Log_Master.xlsx")
YEARS = [2026, 2027, 2028, 2029, 2030]
MONTHS = [
    ("Jan", "January"), ("Feb", "February"), ("Mar", "March"),
    ("Apr", "April"), ("May", "May"), ("Jun", "June"),
    ("Jul", "July"), ("Aug", "August"), ("Sep", "September"),
    ("Oct", "October"), ("Nov", "November"), ("Dec", "December"),
]
LOG_ROWS = 5000
LOG_LAST_ROW = 4 + LOG_ROWS
ALLOWED_FUNCS = {"SUM", "SUMIFS", "COUNTIFS", "COUNTIF", "AVERAGE",
                 "IF", "TEXT", "DATE", "OR", "ABS"}

failures = []


def check(condition, message):
    if not condition:
        failures.append(message)


def month_weeks(year, month):
    first = dt.date(year, month, 1)
    last = dt.date(year, month, calendar.monthrange(year, month)[1])
    start = first - dt.timedelta(days=first.weekday())
    weeks = []
    while start <= last:
        end = start + dt.timedelta(days=6)
        lo, hi = max(start, first), min(end, last)
        if hi >= lo:
            weeks.append((lo, hi))
        start = end + dt.timedelta(days=1)
    return weeks


def month_sheet(year, tab):
    return f"{year} {tab}"


def log_sheet(year):
    return f"{year} Log"


with zipfile.ZipFile(PATH) as z:
    check(z.testzip() is None, "corrupt zip member")
    check("xl/workbook.xml" in z.namelist(), "workbook.xml missing")

wb = load_workbook(PATH)
expected = ["All Years"]
for year in YEARS:
    expected.extend([log_sheet(year), f"{year} Summary"])
    expected.extend(month_sheet(year, tab) for tab, _ in MONTHS)
expected.extend(["Audit Log", "Settings"])
check(wb.sheetnames == expected, f"sheet order wrong ({len(wb.sheetnames)} sheets)")
check(len(wb.sheetnames) == 14 * len(YEARS) + 3, "sheet count")
check(wb.calculation.fullCalcOnLoad is True, "fullCalcOnLoad not set")
check(wb.active.title == "All Years", f"active sheet {wb.active.title}")
check(len(wb.defined_names) == 0, f"named ranges present: {list(wb.defined_names)}")
check(len(wb._external_links) == 0, "external links present")

summary_info = {}

for year in YEARS:
    log = wb[log_sheet(year)]
    check(log.max_row == LOG_LAST_ROW, f"{log.title}: max_row {log.max_row}")
    check(log.max_column == 7, f"{log.title}: max_column {log.max_column}")
    check(log.cell(row=4, column=1).value == "Date", f"{log.title}: header A")
    check(log.cell(row=4, column=2).value == "Day", f"{log.title}: header B")
    check(log.cell(row=4, column=3).value == "Purchaser", f"{log.title}: header C")
    check(log.cell(row=4, column=5).value == "Car Number", f"{log.title}: header E")
    check(log.cell(row=4, column=6).value == "Entered by", f"{log.title}: header F")
    check(log.cell(row=4, column=7).value == "Entered at", f"{log.title}: header G")
    check(log.cell(row=5, column=2).value ==
          '=IF($A5="","",TEXT($A5,"dddd"))', f"{log.title}: day formula")
    check(log.cell(row=LOG_LAST_ROW, column=2).value ==
          f'=IF($A{LOG_LAST_ROW}="","",TEXT($A{LOG_LAST_ROW},"dddd"))',
          f"{log.title}: last day formula")
    check(log.cell(row=5, column=1).value is None, f"{log.title}: should start empty")
    check(log.cell(row=5, column=4).number_format == '"GH₵"#,##0.00',
          f"{log.title}: amount format")
    check(log.cell(row=5, column=7).number_format == "dd mmm yyyy hh:mm",
          f"{log.title}: timestamp format")

    check(log.cell(row=5, column=2).protection.locked is True,
          f"{log.title}: day formula unlocked")
    for col in (1, 3, 4, 5, 6, 7):
        check(log.cell(row=5, column=col).protection.locked is False,
              f"{log.title}: entry col {col} locked")
    check(log.cell(row=LOG_LAST_ROW, column=3).protection.locked is False,
          f"{log.title}: last spare row locked")

    dvs = {dv.type: dv for dv in log.data_validations.dataValidation}
    check(set(dvs) == {"list", "decimal", "date"}, f"{log.title}: DV types {set(dvs)}")
    check(dvs["list"].formula1 == "Settings!$A$15:$A$64",
          f"{log.title}: purchaser DV {dvs['list'].formula1}")
    for rng in dvs["list"].sqref.ranges:
        check(str(rng) == f"C5:C{LOG_LAST_ROW}", f"{log.title}: purchaser DV range {rng}")
    check(dvs["date"].formula1 == str((dt.date(year, 1, 1) - dt.date(1899, 12, 30)).days),
          f"{log.title}: date DV start")
    check(dvs["date"].formula2 == str((dt.date(year, 12, 31) - dt.date(1899, 12, 30)).days),
          f"{log.title}: date DV end")

    cf_ranges = {str(cf.sqref): cf for cf in log.conditional_formatting}
    check(f"A5:E{LOG_LAST_ROW}" in cf_ranges, f"{log.title}: missing row highlights")
    check(len(cf_ranges[f"A5:E{LOG_LAST_ROW}"].rules) == 3,
          f"{log.title}: expected three log highlights")
    check(log.column_dimensions["E"].width == 22, f"{log.title}: car-number width")

    check(log.freeze_panes == "A5", f"{log.title}: freeze {log.freeze_panes}")
    check(log.auto_filter.ref == f"A4:G{LOG_LAST_ROW}",
          f"{log.title}: filter {log.auto_filter.ref}")
    check(log.print_area == f"'{log.title}'!$A$1:$G$4", f"{log.title}: print {log.print_area}")
    check(log.print_title_rows.replace("$", "") == "1:4",
          f"{log.title}: print titles {log.print_title_rows}")
    check(log.page_setup.orientation == "landscape", f"{log.title}: orientation")
    check(log.protection.sheet is True, f"{log.title}: not protected")
    check(log.protection.autoFilter is False, f"{log.title}: filter locked")
    check(log.sheet_properties.tabColor.rgb.endswith("1E4D3D"),
          f"{log.title}: tab colour")

    month_cells = {}
    for m, (tab, full) in enumerate(MONTHS, start=1):
        ws = wb[month_sheet(year, tab)]
        weeks = month_weeks(year, m)
        n = len(weeks)
        week_first, week_last = 6, 5 + n
        total_row = week_last + 1
        count_row = total_row + 1
        avg_row = count_row + 1
        link_row = avg_row + 2
        month_cells[tab] = {"total_row": total_row, "count_row": count_row}

        check(ws.cell(row=5, column=1).value == "Week", f"{ws.title}: header")
        check(ws.cell(row=4, column=1).value == "WEEKLY TOTALS  ·  MON–SUN",
              f"{ws.title}: band")
        ranges = []
        for i in range(n):
            row = week_first + i
            lo = ws.cell(row=row, column=2).value
            hi = ws.cell(row=row, column=3).value
            check(isinstance(lo, dt.datetime) and isinstance(hi, dt.datetime),
                  f"{ws.title}: week {i + 1} dates")
            if isinstance(lo, dt.datetime) and isinstance(hi, dt.datetime):
                ranges.append((lo.date(), hi.date()))
            check(ws.cell(row=row, column=4).value.startswith("=COUNTIFS("),
                  f"{ws.title}: week {i + 1} count")
            check(ws.cell(row=row, column=5).value.startswith("=SUMIFS("),
                  f"{ws.title}: week {i + 1} total")
            check(f"'{log_sheet(year)}'!$A$5:$A${LOG_LAST_ROW}" in
                  ws.cell(row=row, column=5).value,
                  f"{ws.title}: week {i + 1} log reference")
        flattened = []
        for lo, hi in ranges:
            flattened += [lo + dt.timedelta(days=x) for x in range((hi - lo).days + 1)]
        check(flattened == [dt.date(year, m, d)
                            for d in range(1, calendar.monthrange(year, m)[1] + 1)],
              f"{ws.title}: weeks do not match month")

        check(ws.cell(row=total_row, column=5).value ==
              f"=SUM($E$6:$E${week_last})", f"{ws.title}: month total")
        check(ws.cell(row=count_row, column=5).value ==
              f"=SUM($D$6:$D${week_last})", f"{ws.title}: month count")
        check(ws.cell(row=avg_row, column=5).value ==
              f"=IF($E${count_row}=0,0,$E${total_row}/$E${count_row})",
              f"{ws.title}: average")
        link = ws.cell(row=link_row, column=1).hyperlink
        check(link is not None and link.location == f"'{log_sheet(year)}'!A1",
              f"{ws.title}: log hyperlink")
        check(ws.protection.sheet is True, f"{ws.title}: not protected")
        check(ws.print_area == f"'{ws.title}'!$A$1:$E${link_row}",
              f"{ws.title}: print area {ws.print_area}")
        check(ws.print_title_rows.replace("$", "") == "1:5",
              f"{ws.title}: print titles {ws.print_title_rows}")
        check(ws.page_setup.orientation == "portrait", f"{ws.title}: orientation")
        check(ws.sheet_properties.tabColor.rgb.endswith("2F6B57"),
              f"{ws.title}: tab colour")
        check(ws.cell(row=week_first, column=1).protection.locked is True,
              f"{ws.title}: summary cell unlocked")

    summary = wb[f"{year} Summary"]
    for i, (tab, full) in enumerate(MONTHS):
        row = 6 + i
        info = month_cells[tab]
        check(summary.cell(row=row, column=1).value == full, f"{year} Summary month {row}")
        check(summary.cell(row=row, column=2).value ==
              f"='{month_sheet(year, tab)}'!$E${info['count_row']}",
              f"{year} Summary count link {tab}")
        check(summary.cell(row=row, column=3).value ==
              f"='{month_sheet(year, tab)}'!$E${info['total_row']}",
              f"{year} Summary total link {tab}")
        check(summary.cell(row=row, column=4).value == f"=IF($C$18=0,0,C{row}/$C$18)",
              f"{year} Summary share {tab}")
    check(summary.cell(row=18, column=3).value == "=SUM(C6:C17)", f"{year} total")
    check(summary.cell(row=18, column=2).value == "=SUM(B6:B17)", f"{year} count")

    sums = summary.cell(row=22, column=3).value
    counts = summary.cell(row=22, column=2).value
    check("SUMIFS(" in sums and f"'{log_sheet(year)}'!$D$5:$D${LOG_LAST_ROW}" in sums,
          f"{year} purchaser total log reference")
    check(f"DATE({year},1,1)" in sums and f"DATE({year},12,31)" in sums,
          f"{year} purchaser date bounds")
    check("COUNTIFS(" in counts and f"'{log_sheet(year)}'!$C$5:$C${LOG_LAST_ROW}" in counts,
          f"{year} purchaser count log reference")
    check(summary.cell(row=72, column=3).value == "=SUM(C22:C71)",
          f"{year} listed total")
    check(summary.cell(row=73, column=3).value == "=C$18-C72", f"{year} unlisted")
    check(summary.cell(row=75, column=4).value ==
          f"=SUM('{log_sheet(year)}'!$D$5:$D${LOG_LAST_ROW})-$C$18",
          f"{year} integrity check")
    check(summary.cell(row=77, column=1).value.startswith("Unlisted amounts"),
          f"{year} note")
    check(summary.print_area == f"'{year} Summary'!$A$1:$D$77",
          f"{year} print {summary.print_area}")
    check(summary.protection.sheet is True, f"{year} summary not protected")

# ---------------------------------------------------------------- all years
dash = wb["All Years"]
check(not [r for r in dash.merged_cells.ranges
           if r.min_row == r.max_row and r.min_col == r.max_col],
      "All Years has single-cell merges")
for i, year in enumerate(YEARS):
    row = 6 + i
    check(dash.cell(row=row, column=1).value == year, f"All Years year row {row}")
    check(dash.cell(row=row, column=2).value == f"='{year} Summary'!$B$18",
          f"All Years count link {year}")
    check(dash.cell(row=row, column=3).value == f"='{year} Summary'!$C$18",
          f"All Years total link {year}")
    check(dash.cell(row=row, column=4).value == f"=IF($C$11=0,0,C{row}/$C$11)",
          f"All Years share {year}")
    if i == 0:
        check(dash.cell(row=row, column=5).value == "—", "All Years first change")
    else:
        check(dash.cell(row=row, column=5).value ==
              f'=IF(C{row - 1}=0,"—",(C{row}-C{row - 1})/C{row - 1})',
              f"All Years change {year}")
    check(dash.cell(row=row, column=6).value == f"=IF(B{row}=0,0,C{row}/B{row})",
          f"All Years avg purchase {year}")
    check(dash.cell(row=row, column=7).value == f"=C{row}/12",
          f"All Years avg month {year}")
    check(dash.cell(row=row, column=8).value ==
          f'=COUNTIF(\'{year} Summary\'!$B$6:$B$17,">0")',
          f"All Years months logged {year}")
check(dash.cell(row=11, column=3).value == "=SUM(C6:C10)", "All Years total")
check(dash.cell(row=11, column=6).value == "=IF(B11=0,0,C11/B11)",
      "All Years avg purchase total")
check(dash.cell(row=11, column=7).value == "=AVERAGE(G6:G10)",
      "All Years avg month total")
check(dash.cell(row=11, column=8).value == "=SUM(H6:H10)",
      "All Years months total")

for i in range(50):
    row = 15 + i
    settings_row = 15 + i
    check(dash.cell(row=row, column=1).value ==
          f'=IF(Settings!A{settings_row}="","",Settings!A{settings_row})',
          f"All Years purchaser name row {row}")
    for j, year in enumerate(YEARS):
        expected_formula = (
            f'=IF($A{row}="","",IF(\'{year} Summary\'!$C${22 + i}=0,"",'
            f"'{year} Summary'!$C${22 + i}))"
        )
        check(dash.cell(row=row, column=2 + j).value == expected_formula,
              f"All Years matrix link {year} row {row}")
    check(dash.cell(row=row, column=7).value ==
          f'=IF($A{row}="","",SUM(B{row}:F{row}))', f"All Years row total {row}")
    check(dash.cell(row=row, column=8).value ==
          f'=IF(OR($A{row}="",$C$11=0),0,G{row}/$C$11)',
          f"All Years row pct {row}")
check(dash.cell(row=65, column=1).value == "Listed purchasers", "All Years listed")
check(dash.cell(row=66, column=1).value == "Unlisted / blank purchaser",
      "All Years unlisted")
check(dash.print_area == "'All Years'!$A$1:$H$68", f"All Years print {dash.print_area}")
check(dash.protection.sheet is True, "All Years not protected")

# ---------------------------------------------------------------- audit log
audit = wb["Audit Log"]
audit_headers = ["Timestamp", "User", "Sheet", "Row", "Purchase date", "Purchaser",
                 "Amount (GH₵)", "Car Number", "Field", "Old value",
                 "New value", "Action"]
for i, text in enumerate(audit_headers, start=1):
    check(audit.cell(row=4, column=i).value == text, f"Audit Log header {i}")
check(audit.max_row == 4, f"Audit Log should be empty; max_row={audit.max_row}")
check(audit.freeze_panes == "A5", f"Audit Log freeze {audit.freeze_panes}")
check(str(audit.auto_filter.ref) == "A4:L1000",
      f"Audit Log filter {audit.auto_filter.ref}")
check(audit.protection.sheet is True, "Audit Log not protected")
check(audit.protection.autoFilter is False, "Audit Log filter locked")
check(audit.print_area in (None, ""), f"Audit Log print area {audit.print_area}")

# ---------------------------------------------------------------- settings
settings = wb["Settings"]
check(settings.cell(row=14, column=1).value == "Purchaser Name", "Settings header")
check("2026–2030" in settings.cell(row=2, column=1).value, "Settings year range")
check("Add Purchase" in settings.cell(row=13, column=1).value, "Settings audit note")
check("Audit Log" in settings.cell(row=13, column=1).value, "Settings audit log note")
for row in (15, 64):
    check(settings.cell(row=row, column=1).protection.locked is False,
          f"Settings A{row} should be editable")
check(settings.protection.sheet is True, "Settings not protected")
check(settings.print_area == "'Settings'!$A$1:$B$68",
      f"Settings print {settings.print_area}")
check(settings.cell(row=67, column=1).value == "Mobile form link", "Settings form link label")
check(settings.cell(row=68, column=1).value == "Owner account", "Settings owner label")
check(settings.cell(row=68, column=2).value == "Set by administrator",
      "Settings owner placeholder")

# ---------------------------------------------------------------- formulas
formula_cells = 0
funcs = set()
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_cells += 1
                funcs.update(re.findall(r"([A-Z][A-Z0-9\.]*)\(", cell.value))
unsupported = funcs - ALLOWED_FUNCS
check(not unsupported, f"unsupported functions: {sorted(unsupported)}")
check(formula_cells > 25000, f"unexpectedly few formulas: {formula_cells}")

print(f"sheets: {len(wb.sheetnames)} ({len(YEARS)} years)")
print(f"formulas checked: {formula_cells}")
print(f"functions used: {', '.join(sorted(funcs))}")
print(f"merged ranges: {sum(1 for ws in wb.worksheets for _ in ws.merged_cells.ranges)}")
if failures:
    print(f"\nFAILED ({len(failures)}):")
    for failure in failures:
        print(" -", failure)
    raise SystemExit(1)
print("\nAll checks passed.")
