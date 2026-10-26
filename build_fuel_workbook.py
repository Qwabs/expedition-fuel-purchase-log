#!/usr/bin/env python3
"""Build Fuel_Purchase_Log_Master.xlsx (2026-2030).

Each year has a Summary plus 12 monthly sheets; weekly totals are
Monday-Sunday and clipped to the month. All Years combines the annual
summaries. `add_year.py` appends later years without touching data.
"""

import argparse
import calendar
import datetime as dt
import re
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties

YEARS = [2026, 2027, 2028, 2029, 2030]
OUT = Path(__file__).with_name("Fuel_Purchase_Log_Master.xlsx")
LOG_ROWS = 5000
LOG_LAST_ROW = 4 + LOG_ROWS
SPARE_ROWS = 15

MONTHS = [
    ("Jan", "January"), ("Feb", "February"), ("Mar", "March"),
    ("Apr", "April"), ("May", "May"), ("Jun", "June"),
    ("Jul", "July"), ("Aug", "August"), ("Sep", "September"),
    ("Oct", "October"), ("Nov", "November"), ("Dec", "December"),
]

LIST_FIRST_ROW = 15
LIST_LAST_ROW = 64
PURCHASER_SLOTS = LIST_LAST_ROW - LIST_FIRST_ROW + 1
SUMMARY_FIRST_PURCH_ROW = 22
SUMMARY_UNLISTED_ROW = 73

GREEN = "1E4D3D"
GREEN_SOFT = "2F6B57"
GOLD = "B08D57"
GOLD_LIGHT = "EFE4CF"
CREAM = "FBF7EC"
BAND = "F3F8F4"
SPARE = "FDFCF8"
LINE = "D6DDD8"
GREY = "6B7280"
INK = "374151"
RED_FILL = "FDE2E1"
RED_TEXT = "9C0006"

CURRENCY_FMT = '"GH₵"#,##0.00'
DATE_FMT = "d mmm yyyy"
SHORT_DATE_FMT = "d mmm"
FONT = "Arial"

TITLE_FONT = Font(name=FONT, size=16, bold=True, color=GREEN)
SUBTITLE_FONT = Font(name=FONT, size=9, italic=True, color=GREY)
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10, color=INK)
BODY_BOLD = Font(name=FONT, size=10, bold=True, color=INK)
SMALL_GREY = Font(name=FONT, size=8.5, color=GREY)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color=GREY)

THIN = Side(style="thin", color=LINE)
HAIR = Side(style="hair", color=LINE)
ALL_BORDERS = Border(left=THIN, right=THIN, top=HAIR, bottom=HAIR)
HEADER_BORDER = Border(
    left=THIN, right=THIN,
    top=Side(style="thin", color=GREEN),
    bottom=Side(style="medium", color=GOLD),
)
UNLOCKED = Protection(locked=False)
LOCKED = Protection(locked=True)


def fill(color):
    return PatternFill("solid", fgColor=color)


def month_sheet(year, tab):
    return f"{year} {tab}"


def year_summary_sheet(year):
    return f"{year} Summary"


def log_sheet(year):
    return f"{year} Log"


def month_weeks(year, month):
    first = dt.date(year, month, 1)
    last = dt.date(year, month, calendar.monthrange(year, month)[1])
    start = first - dt.timedelta(days=first.weekday())
    weeks = []
    while start <= last:
        end = start + dt.timedelta(days=6)
        lo, hi = max(start, first), min(end, last)
        if hi >= lo:
            weeks.append((lo, hi))
        start = end + dt.timedelta(days=1)
    return weeks


def week_index(weeks, day):
    for i, (lo, hi) in enumerate(weeks):
        if lo <= day <= hi:
            return i
    raise ValueError(day)


def excel_serial(day):
    return (day - dt.date(1899, 12, 30)).days


def merge_after_styling(ws, cell_range, value=None, font=None, fill_color=None,
                        alignment=None, border=None, number_format=None):
    rows = ws[cell_range]
    for row in rows:
        for cell in row:
            if fill_color:
                cell.fill = fill(fill_color)
            if border:
                cell.border = border
    top_left = rows[0][0]
    if value is not None:
        top_left.value = value
    if font:
        top_left.font = font
    if alignment:
        top_left.alignment = alignment
    if number_format:
        top_left.number_format = number_format
    ws.merge_cells(cell_range)
    return top_left


def detect_years(wb):
    years = []
    for title in wb.sheetnames:
        match = re.fullmatch(r"(\d{4}) Summary", title)
        if match:
            years.append(int(match.group(1)))
    return sorted(set(years))


def clear_sheet(ws):
    ws.protection.sheet = False
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    if ws.max_row > 0:
        ws.delete_rows(1, ws.max_row)
    ws.conditional_formatting = ConditionalFormattingList()
    ws.data_validations.dataValidation = []
    ws.print_area = None
    ws.print_title_rows = None
    ws.freeze_panes = None


def reorder_sheets(wb, years):
    order = ["All Years"]
    for year in years:
        order.append(log_sheet(year))
        order.append(year_summary_sheet(year))
        order.extend(month_sheet(year, tab) for tab, _ in MONTHS)
    order.extend(["Audit Log", "Settings"])
    sheets = {ws.title: ws for ws in wb.worksheets}
    wb._sheets = [sheets[title] for title in order if title in sheets]


# ------------------------------------------------------------------ year log
def build_log(ws, year):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = GREEN

    widths = {"A": 13, "B": 13, "C": 26, "D": 16, "E": 22, "F": 28, "G": 21}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 30
    merge_after_styling(
        ws, "A1:G1",
        value=f"FUEL PURCHASE LOG   ·   {year}",
        font=TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=Side(style="medium", color=GOLD)),
    )
    ws.row_dimensions[2].height = 20
    merge_after_styling(
        ws, "A2:G2",
        value=("One row per purchase  ·  Use Fuel Log → Add Purchase or the next blank row  ·  "
               "Multiple purchases on the same day are expected"),
        font=SUBTITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[3].height = 20
    merge_after_styling(
        ws, "A3:G3",
        value=("Date, purchaser, amount and car number are entered here. Day and audit columns "
               "fill automatically. Monthly summaries read from this log."),
        font=Font(name=FONT, size=9, italic=True, color=GREY),
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )

    headers = ["Date", "Day", "Purchaser", "Amount (GH₵)", "Car Number",
               "Entered by", "Entered at"]
    ws.row_dimensions[4].height = 30
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=text)
        c.fill = fill("4B5563" if i >= 6 else GREEN)
        c.font = HEADER_FONT
        c.border = HEADER_BORDER
        c.alignment = Alignment(
            horizontal="center" if i in (1, 2, 4) else "left",
            vertical="center", indent=0 if i in (1, 2, 4) else 1,
        )

    row_border = Border(left=HAIR, right=HAIR, bottom=HAIR)
    audit_fill = fill("F7F7F7")
    for row in range(5, LOG_LAST_ROW + 1):
        for col in range(1, 8):
            c = ws.cell(row=row, column=col)
            c.border = row_border
            c.font = BODY_FONT
            c.alignment = Alignment(
                horizontal="center" if col == 2 else
                ("right" if col == 4 else "left"),
                vertical="center", indent=0 if col in (2,) else 1,
            )
            if col == 1:
                c.number_format = DATE_FMT
                c.protection = UNLOCKED
            elif col == 2:
                c.value = f'=IF($A{row}="","",TEXT($A{row},"dddd"))'
                c.font = Font(name=FONT, size=10, color=GREY)
                c.protection = LOCKED
            elif col == 3:
                c.protection = UNLOCKED
            elif col == 4:
                c.number_format = CURRENCY_FMT
                c.protection = UNLOCKED
            elif col == 5:
                c.font = Font(name=FONT, size=9.5, color="4B5563")
                c.protection = UNLOCKED
            elif col == 6:
                c.fill = audit_fill
                c.font = Font(name=FONT, size=9, color="4B5563")
                c.protection = UNLOCKED
            else:
                c.fill = audit_fill
                c.number_format = "dd mmm yyyy hh:mm"
                c.font = Font(name=FONT, size=9, color="4B5563")
                c.protection = UNLOCKED

    dv_purchaser = DataValidation(
        type="list",
        formula1=f"Settings!$A${LIST_FIRST_ROW}:$A${LIST_LAST_ROW}",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="warning",
        errorTitle="Not in the purchaser list",
        error=("This name is not in Settings!A15:A64. Add it there so it appears in the "
               "dropdown; click Yes to keep it for now."),
        showInputMessage=True,
        promptTitle="Purchaser",
        prompt="Pick a name from Settings, or type a new one.",
    )
    ws.add_data_validation(dv_purchaser)
    dv_purchaser.add(f"C5:C{LOG_LAST_ROW}")

    dv_amount = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="warning",
        errorTitle="Check the amount",
        error="Amounts are normally zero or more. Click Yes to keep this value.",
    )
    ws.add_data_validation(dv_amount)
    dv_amount.add(f"D5:D{LOG_LAST_ROW}")

    first = dt.date(year, 1, 1)
    last = dt.date(year, 12, 31)
    dv_date = DataValidation(
        type="date",
        operator="between",
        formula1=str(excel_serial(first)),
        formula2=str(excel_serial(last)),
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="warning",
        errorTitle="Date outside this year",
        error=(f"This date is outside {year}. Click Yes if intentional; it will not be "
               "counted in the year totals."),
    )
    ws.add_data_validation(dv_date)
    dv_date.add(f"A5:A{LOG_LAST_ROW}")

    red_fill = fill(RED_FILL)
    ws.conditional_formatting.add(
        f"A5:E{LOG_LAST_ROW}",
        FormulaRule(formula=['AND($D5<>"",$A5="")'], fill=red_fill,
                    font=Font(name=FONT, size=10, color=RED_TEXT)),
    )
    ws.conditional_formatting.add(
        f"A5:E{LOG_LAST_ROW}",
        FormulaRule(
            formula=[f'AND($A5<>"",OR($A5<DATE({year},1,1),$A5>DATE({year},12,31)))'],
            fill=fill("FEF3C7"),
            font=Font(name=FONT, size=10, color="92400E"),
        ),
    )
    ws.conditional_formatting.add(
        f"A5:E{LOG_LAST_ROW}",
        FormulaRule(
            formula=['AND($A5<>"",$D5<>"",$E5="")'],
            fill=fill("FEF3C7"),
            font=Font(name=FONT, size=10, color="92400E"),
        ),
    )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:G{LOG_LAST_ROW}"

    ws.print_area = "A1:G4"
    ws.print_title_rows = "1:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"

    ws.protection.sheet = True
    ws.protection.autoFilter = False


# ------------------------------------------------------- monthly summary
def build_month_summary(ws, year, month_no, month_full):
    days = calendar.monthrange(year, month_no)[1]
    first = dt.date(year, month_no, 1)
    last = dt.date(year, month_no, days)
    weeks = month_weeks(year, month_no)
    n_weeks = len(weeks)

    week_first = 6
    week_last = week_first + n_weeks - 1
    total_row = week_last + 1
    count_row = total_row + 1
    avg_row = count_row + 1
    link_row = avg_row + 2

    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_SOFT
    widths = {"A": 10, "B": 12, "C": 12, "D": 11, "E": 18}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 32
    merge_after_styling(
        ws, "A1:E1",
        value=f"{month_full.upper()} {year}   ·   WEEKLY SUMMARY",
        font=TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=Side(style="medium", color=GOLD)),
    )
    ws.row_dimensions[2].height = 22
    merge_after_styling(
        ws, "A2:E2",
        value=(f"Read-only summary from the {year} Log  ·  "
               "Weekly totals run Monday to Sunday"),
        font=SUBTITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[3].height = 10

    ws.row_dimensions[4].height = 26
    merge_after_styling(
        ws, "A4:E4",
        value="WEEKLY TOTALS  ·  MON–SUN",
        font=HEADER_FONT,
        fill_color=GREEN,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
    )

    headers = ["Week", "From", "To", "Purchases", "Total (GH₵)"]
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=5, column=i, value=text)
        c.fill = fill(GOLD_LIGHT)
        c.font = Font(name=FONT, size=9, bold=True, color="6B4F2A")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="center", vertical="center")

    log_ref = f"'{log_sheet(year)}'"
    for i, (lo, hi) in enumerate(weeks):
        row = week_first + i
        ws.row_dimensions[row].height = 21
        band = "FFFFFF" if i % 2 == 0 else "FBF9F2"
        for col in range(1, 6):
            c = ws.cell(row=row, column=col)
            c.fill = fill(band)
            c.border = ALL_BORDERS
            c.font = BODY_FONT
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.protection = LOCKED
        ws.cell(row=row, column=1, value=f"Wk {i + 1}").font = BODY_BOLD
        ws.cell(row=row, column=2, value=lo).number_format = SHORT_DATE_FMT
        ws.cell(row=row, column=3, value=hi).number_format = SHORT_DATE_FMT
        count = ws.cell(
            row=row, column=4,
            value=(f'=COUNTIFS({log_ref}!$A$5:$A${LOG_LAST_ROW},">="&B{row},'
                   f'{log_ref}!$A$5:$A${LOG_LAST_ROW},"<="&C{row},'
                   f'{log_ref}!$D$5:$D${LOG_LAST_ROW},">0")'),
        )
        total = ws.cell(
            row=row, column=5,
            value=(f'=SUMIFS({log_ref}!$D$5:$D${LOG_LAST_ROW},'
                   f'{log_ref}!$A$5:$A${LOG_LAST_ROW},">="&B{row},'
                   f'{log_ref}!$A$5:$A${LOG_LAST_ROW},"<="&C{row})'),
        )
        count.number_format = "0"
        count.alignment = Alignment(horizontal="center", vertical="center")
        total.number_format = CURRENCY_FMT
        total.font = BODY_BOLD
        total.alignment = Alignment(horizontal="right", vertical="center", indent=1)

    ws.row_dimensions[total_row].height = 26
    merge_after_styling(
        ws, f"A{total_row}:D{total_row}",
        value="MONTH TOTAL",
        font=HEADER_FONT,
        fill_color=GREEN,
        alignment=Alignment(horizontal="right", vertical="center", indent=1),
        border=Border(left=THIN, top=THIN, bottom=THIN),
    )
    total_cell = ws.cell(row=total_row, column=5,
                         value=f"=SUM($E${week_first}:$E${week_last})")
    total_cell.fill = fill(GREEN)
    total_cell.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    total_cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    total_cell.number_format = CURRENCY_FMT
    total_cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    total_cell.protection = LOCKED

    ws.row_dimensions[count_row].height = 20
    merge_after_styling(
        ws, f"A{count_row}:D{count_row}",
        value="Purchases logged",
        font=Font(name=FONT, size=9, color=INK),
        fill_color=CREAM,
        alignment=Alignment(horizontal="right", vertical="center", indent=1),
        border=Border(left=THIN, bottom=THIN),
    )
    count_cell = ws.cell(row=count_row, column=5,
                         value=f"=SUM($D${week_first}:$D${week_last})")
    count_cell.fill = fill(CREAM)
    count_cell.font = Font(name=FONT, size=9, bold=True, color=INK)
    count_cell.border = Border(left=THIN, right=THIN, bottom=THIN)
    count_cell.number_format = "0"
    count_cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    count_cell.protection = LOCKED

    ws.row_dimensions[avg_row].height = 20
    merge_after_styling(
        ws, f"A{avg_row}:D{avg_row}",
        value="Average purchase",
        font=Font(name=FONT, size=9, italic=True, color=GREY),
        alignment=Alignment(horizontal="right", vertical="center", indent=1),
        border=Border(left=THIN, bottom=THIN),
    )
    avg_cell = ws.cell(
        row=avg_row, column=5,
        value=f'=IF($E${count_row}=0,0,$E${total_row}/$E${count_row})',
    )
    avg_cell.font = Font(name=FONT, size=9, color=GREY)
    avg_cell.border = Border(left=THIN, right=THIN, bottom=THIN)
    avg_cell.number_format = CURRENCY_FMT
    avg_cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    avg_cell.protection = LOCKED

    ws.row_dimensions[link_row].height = 24
    link = merge_after_styling(
        ws, f"A{link_row}:E{link_row}",
        value=f"Open the {year} Log to add or correct purchases",
        font=Font(name=FONT, size=9, bold=True, color=GREEN, underline="single"),
        fill_color="F7FAF8",
        alignment=Alignment(horizontal="center", vertical="center"),
        border=Border(top=THIN, bottom=THIN),
    )
    link.hyperlink = Hyperlink(
        ref=link.coordinate, location=f"'{log_sheet(year)}'!A1",
        display=f"Open the {year} Log",
    )
    link.protection = LOCKED

    ws.print_area = f"A1:E{link_row}"
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"

    ws.protection.sheet = True
    return {"total_row": total_row, "count_row": count_row}


# ---------------------------------------------------------------- year summary
def build_year_summary(ws, year, month_cells):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GOLD
    widths = {"A": 22, "B": 13, "C": 20, "D": 13}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 34
    merge_after_styling(
        ws, "A1:D1",
        value=f"FUEL PURCHASE SUMMARY   ·   {year}",
        font=TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=Side(style="medium", color=GOLD)),
    )
    ws.row_dimensions[2].height = 22
    merge_after_styling(
        ws, "A2:D2",
        value=(f"Calculated from the {year} Log  ·  Ghana Cedi (GH₵)  ·  "
               "Weeks Monday to Sunday"),
        font=SUBTITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[3].height = 10

    ws.row_dimensions[4].height = 26
    merge_after_styling(
        ws, "A4:D4",
        value="MONTHLY TOTALS",
        font=HEADER_FONT,
        fill_color=GREEN,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
    )

    headers = ["Month", "Purchases", "Total (GH₵)", "% of year"]
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=5, column=i, value=text)
        c.fill = fill(GOLD_LIGHT)
        c.font = Font(name=FONT, size=9, bold=True, color="6B4F2A")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, (tab, full) in enumerate(MONTHS):
        row = 6 + i
        ws.row_dimensions[row].height = 21
        band = "FFFFFF" if i % 2 == 0 else "FBF9F2"
        month_cell = ws.cell(row=row, column=1, value=full)
        month_cell.fill = fill(band)
        month_cell.font = Font(name=FONT, size=10, bold=True, color=GREEN, underline="single")
        month_cell.border = ALL_BORDERS
        month_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        month_cell.hyperlink = Hyperlink(
            ref=month_cell.coordinate,
            location=f"'{month_sheet(year, tab)}'!A1",
            display=full,
        )
        month_cell.protection = LOCKED

        info = month_cells[tab]
        count = ws.cell(row=row, column=2,
                        value=f"='{month_sheet(year, tab)}'!$E${info['count_row']}")
        total = ws.cell(row=row, column=3,
                        value=f"='{month_sheet(year, tab)}'!$E${info['total_row']}")
        share = ws.cell(row=row, column=4,
                        value=f"=IF($C$18=0,0,C{row}/$C$18)")
        for c in (count, total, share):
            c.fill = fill(band)
            c.font = BODY_FONT
            c.border = ALL_BORDERS
            c.protection = LOCKED
        count.number_format = "0"
        count.alignment = Alignment(horizontal="center", vertical="center")
        total.number_format = CURRENCY_FMT
        total.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        share.number_format = "0.0%"
        share.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[18].height = 25
    year_label = ws.cell(row=18, column=1, value="YEAR TOTAL")
    year_label.font = HEADER_FONT
    year_label.fill = fill(GREEN)
    year_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    year_label.border = Border(left=THIN, top=THIN, bottom=THIN)
    year_label.protection = LOCKED

    year_count = ws.cell(row=18, column=2, value="=SUM(B6:B17)")
    year_total = ws.cell(row=18, column=3, value="=SUM(C6:C17)")
    year_share = ws.cell(row=18, column=4, value="=SUM(D6:D17)")
    for col in range(1, 5):
        c = ws.cell(row=18, column=col)
        c.fill = fill(GREEN)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        c.border = Border(left=THIN, right=THIN,
                          top=Side(style="thin", color=GREEN),
                          bottom=Side(style="double", color=GREEN))
        c.protection = LOCKED
    year_count.number_format = "0"
    year_count.alignment = Alignment(horizontal="center", vertical="center")
    year_total.number_format = CURRENCY_FMT
    year_total.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    year_share.number_format = "0.0%"
    year_share.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[19].height = 22
    log_link = merge_after_styling(
        ws, "A19:D19",
        value=f"Open the {year} Log to add or correct purchases",
        font=Font(name=FONT, size=9, bold=True, color=GREEN, underline="single"),
        fill_color="F7FAF8",
        alignment=Alignment(horizontal="center", vertical="center"),
        border=Border(top=THIN, bottom=THIN),
    )
    log_link.hyperlink = Hyperlink(
        ref=log_link.coordinate, location=f"'{log_sheet(year)}'!A1",
        display=f"Open the {year} Log",
    )
    log_link.protection = LOCKED

    ws.row_dimensions[20].height = 26
    merge_after_styling(
        ws, "A20:D20",
        value="SPEND BY PURCHASER",
        font=HEADER_FONT,
        fill_color=GREEN,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
    )

    purch_headers = ["Purchaser", "Purchases", "Total (GH₵)", "% of year"]
    for i, text in enumerate(purch_headers, start=1):
        c = ws.cell(row=21, column=i, value=text)
        c.fill = fill(GOLD_LIGHT)
        c.font = Font(name=FONT, size=9, bold=True, color="6B4F2A")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="center", vertical="center")

    log_ref = f"'{log_sheet(year)}'"
    sums = (
        f'SUMIFS({log_ref}!$D$5:$D${LOG_LAST_ROW},'
        f'{log_ref}!$C$5:$C${LOG_LAST_ROW},{{ref}},'
        f'{log_ref}!$A$5:$A${LOG_LAST_ROW},">="&DATE({year},1,1),'
        f'{log_ref}!$A$5:$A${LOG_LAST_ROW},"<="&DATE({year},12,31))'
    )
    counts = (
        f'COUNTIFS({log_ref}!$C$5:$C${LOG_LAST_ROW},{{ref}},'
        f'{log_ref}!$D$5:$D${LOG_LAST_ROW},">0",'
        f'{log_ref}!$A$5:$A${LOG_LAST_ROW},">="&DATE({year},1,1),'
        f'{log_ref}!$A$5:$A${LOG_LAST_ROW},"<="&DATE({year},12,31))'
    )

    for i in range(PURCHASER_SLOTS):
        row = SUMMARY_FIRST_PURCH_ROW + i
        settings_row = LIST_FIRST_ROW + i
        ws.row_dimensions[row].height = 20
        band = "FFFFFF" if i % 2 == 0 else "FBF9F2"

        name = ws.cell(
            row=row, column=1,
            value=f'=IF(Settings!A{settings_row}="","",Settings!A{settings_row})',
        )
        name.fill = fill(band)
        name.font = BODY_FONT
        name.border = ALL_BORDERS
        name.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        name.protection = LOCKED

        count = ws.cell(
            row=row, column=2,
            value=f'=IF($A{row}="","",' + counts.format(ref=f"$A{row}") + ")",
        )
        total = ws.cell(
            row=row, column=3,
            value=f'=IF($A{row}="","",' + sums.format(ref=f"$A{row}") + ")",
        )
        share = ws.cell(
            row=row, column=4,
            value=f'=IF(OR($A{row}="",$C$18=0),0,C{row}/$C$18)',
        )
        for c in (count, total, share):
            c.fill = fill(band)
            c.font = BODY_FONT
            c.border = ALL_BORDERS
            c.protection = LOCKED
        count.number_format = "0"
        count.alignment = Alignment(horizontal="center", vertical="center")
        total.number_format = CURRENCY_FMT
        total.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        share.number_format = "0.0%"
        share.alignment = Alignment(horizontal="center", vertical="center")

    listed_row = SUMMARY_UNLISTED_ROW - 1
    ws.row_dimensions[listed_row].height = 24
    listed_label = ws.cell(row=listed_row, column=1, value="Listed purchasers")
    listed_label.font = BODY_BOLD
    listed_label.fill = fill(CREAM)
    listed_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    listed_label.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    listed_label.protection = LOCKED
    ws.cell(row=listed_row, column=2,
            value=f"=SUM(B{SUMMARY_FIRST_PURCH_ROW}:B{SUMMARY_UNLISTED_ROW - 2})")
    ws.cell(row=listed_row, column=3,
            value=f"=SUM(C{SUMMARY_FIRST_PURCH_ROW}:C{SUMMARY_UNLISTED_ROW - 2})")
    ws.cell(row=listed_row, column=4, value=f"=IF($C$18=0,0,C{listed_row}/$C$18)")
    for col in range(1, 5):
        c = ws.cell(row=listed_row, column=col)
        c.fill = fill(CREAM)
        c.font = BODY_BOLD
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        c.protection = LOCKED
    ws.cell(row=listed_row, column=2).number_format = "0"
    ws.cell(row=listed_row, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=listed_row, column=3).number_format = CURRENCY_FMT
    ws.cell(row=listed_row, column=3).alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.cell(row=listed_row, column=4).number_format = "0.0%"
    ws.cell(row=listed_row, column=4).alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[SUMMARY_UNLISTED_ROW].height = 22
    unlisted_label = ws.cell(row=SUMMARY_UNLISTED_ROW, column=1,
                             value="Unlisted / blank purchaser")
    unlisted_label.font = Font(name=FONT, size=9, italic=True, color=GREY)
    unlisted_label.fill = fill("FFFFFF")
    unlisted_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    unlisted_label.border = Border(left=THIN, right=THIN, bottom=THIN)
    unlisted_label.protection = LOCKED
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=2, value=f"=B$18-B{listed_row}")
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=3, value=f"=C$18-C{listed_row}")
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=4,
            value=f"=IF($C$18=0,0,C{SUMMARY_UNLISTED_ROW}/$C$18)")
    for col in range(1, 5):
        c = ws.cell(row=SUMMARY_UNLISTED_ROW, column=col)
        c.fill = fill("FFFFFF")
        c.font = Font(name=FONT, size=9, italic=True, color=GREY)
        c.border = Border(left=THIN, right=THIN, bottom=THIN)
        c.protection = LOCKED
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=2).number_format = "0"
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=3).number_format = CURRENCY_FMT
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=3).alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=4).number_format = "0.0%"
    ws.cell(row=SUMMARY_UNLISTED_ROW, column=4).alignment = Alignment(horizontal="center", vertical="center")

    ws.conditional_formatting.add(
        f"A{SUMMARY_UNLISTED_ROW}:D{SUMMARY_UNLISTED_ROW}",
        FormulaRule(
            formula=[f"ABS($C${SUMMARY_UNLISTED_ROW})>0.005"],
            fill=fill(RED_FILL),
            font=Font(name=FONT, size=9, bold=True, color=RED_TEXT),
        ),
    )

    check_row = SUMMARY_UNLISTED_ROW + 2
    ws.row_dimensions[check_row].height = 22
    merge_after_styling(
        ws, f"A{check_row}:C{check_row}",
        value=f"Unposted / outside {year} (should be 0.00)",
        font=Font(name=FONT, size=8.5, italic=True, color=GREY),
        alignment=Alignment(horizontal="right", vertical="center", indent=1),
        border=Border(top=HAIR, bottom=THIN),
    )
    check_cell = ws.cell(
        row=check_row, column=4,
        value=f"=SUM({log_ref}!$D$5:$D${LOG_LAST_ROW})-$C$18",
    )
    check_cell.font = Font(name=FONT, size=9, color=GREY)
    check_cell.border = Border(left=THIN, right=THIN, top=HAIR, bottom=THIN)
    check_cell.number_format = CURRENCY_FMT
    check_cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    check_cell.protection = LOCKED
    ws.conditional_formatting.add(
        f"D{check_row}",
        FormulaRule(
            formula=[f"ABS($D${check_row})>0.005"],
            fill=fill(RED_FILL),
            font=Font(name=FONT, size=9, bold=True, color=RED_TEXT),
        ),
    )

    note_row = check_row + 2
    ws.row_dimensions[note_row].height = 34
    merge_after_styling(
        ws, f"A{note_row}:D{note_row}",
        value=("Unlisted amounts are included in the year total. Add the purchaser's name in "
               "Settings to itemise them. The check line flags missing or out-of-year dates."),
        font=NOTE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1),
        border=Border(top=HAIR),
    )

    ws.print_area = f"A1:D{note_row}"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"

    ws.protection.sheet = True


# ---------------------------------------------------------------- all years
def build_all_years(ws, years):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN
    ws.column_dimensions["A"].width = 26
    last_year_col = 1 + len(years)
    total_col = last_year_col + 1
    pct_col = total_col + 1
    for col in range(2, pct_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.column_dimensions[get_column_letter(pct_col)].width = 14

    last_col_letter = get_column_letter(pct_col)

    ws.row_dimensions[1].height = 34
    merge_after_styling(
        ws, f"A1:{last_col_letter}1",
        value="FUEL PURCHASE SUMMARY   ·   ALL YEARS",
        font=TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=Side(style="medium", color=GOLD)),
    )
    ws.row_dimensions[2].height = 22
    merge_after_styling(
        ws, f"A2:{last_col_letter}2",
        value=(f"{min(years)}–{max(years)}  ·  Ghana Cedi (GH₵)  ·  Weeks Monday to Sunday  ·  "
               "All figures formula-linked"),
        font=SUBTITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[3].height = 10

    ws.row_dimensions[4].height = 26
    merge_after_styling(
        ws, f"A4:{last_col_letter}4",
        value="YEARLY TOTALS",
        font=HEADER_FONT,
        fill_color=GREEN,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
    )

    headers = [
        "Year", "Purchases", "Total (GH₵)", "% of all time",
        "Change vs prior year", "Avg / purchase", "Avg / month", "Months logged",
    ]
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=5, column=i, value=text)
        c.fill = fill(GOLD_LIGHT)
        c.font = Font(name=FONT, size=9, bold=True, color="6B4F2A")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="center", vertical="center")

    first_year_row = 6
    last_year_row = first_year_row + len(years) - 1
    year_total_row = last_year_row + 1
    for i, year in enumerate(years):
        row = first_year_row + i
        ws.row_dimensions[row].height = 22
        band = "FFFFFF" if i % 2 == 0 else "FBF9F2"

        year_cell = ws.cell(row=row, column=1, value=year)
        year_cell.number_format = "0"
        year_cell.fill = fill(band)
        year_cell.font = Font(name=FONT, size=10, bold=True, color=GREEN, underline="single")
        year_cell.border = ALL_BORDERS
        year_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        year_cell.hyperlink = Hyperlink(
            ref=year_cell.coordinate,
            location=f"'{year_summary_sheet(year)}'!A1",
            display=str(year),
        )
        year_cell.protection = LOCKED

        count = ws.cell(row=row, column=2,
                        value=f"='{year_summary_sheet(year)}'!$B$18")
        total = ws.cell(row=row, column=3,
                        value=f"='{year_summary_sheet(year)}'!$C$18")
        share = ws.cell(row=row, column=4,
                        value=f"=IF($C${year_total_row}=0,0,C{row}/$C${year_total_row})")
        if i == 0:
            change = ws.cell(row=row, column=5, value="—")
        else:
            change = ws.cell(
                row=row, column=5,
                value=f'=IF(C{row - 1}=0,"—",(C{row}-C{row - 1})/C{row - 1})',
            )
        avg_purchase = ws.cell(
            row=row, column=6,
            value=f"=IF(B{row}=0,0,C{row}/B{row})",
        )
        avg_month = ws.cell(row=row, column=7, value=f"=C{row}/12")
        months_logged = ws.cell(
            row=row, column=8,
            value=(f'=COUNTIF(\'{year_summary_sheet(year)}\'!$B$6:$B$17,">0")'),
        )
        for c in (count, total, share, change, avg_purchase, avg_month, months_logged):
            c.fill = fill(band)
            c.font = BODY_FONT
            c.border = ALL_BORDERS
            c.protection = LOCKED
        count.number_format = "0"
        count.alignment = Alignment(horizontal="center", vertical="center")
        total.number_format = CURRENCY_FMT
        total.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        share.number_format = "0.0%"
        share.alignment = Alignment(horizontal="center", vertical="center")
        change.number_format = "0.0%;-0.0%"
        change.alignment = Alignment(horizontal="center", vertical="center")
        avg_purchase.number_format = CURRENCY_FMT
        avg_purchase.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        avg_month.number_format = CURRENCY_FMT
        avg_month.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        months_logged.number_format = "0"
        months_logged.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[year_total_row].height = 25
    total_label = ws.cell(row=year_total_row, column=1, value="ALL YEARS")
    total_label.font = HEADER_FONT
    total_label.fill = fill(GREEN)
    total_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    total_label.border = Border(left=THIN, top=THIN, bottom=THIN)
    total_label.protection = LOCKED
    all_count = ws.cell(row=year_total_row, column=2,
                        value=f"=SUM(B{first_year_row}:B{last_year_row})")
    all_total = ws.cell(row=year_total_row, column=3,
                        value=f"=SUM(C{first_year_row}:C{last_year_row})")
    all_share = ws.cell(row=year_total_row, column=4,
                        value=f"=SUM(D{first_year_row}:D{last_year_row})")
    all_avg_purchase = ws.cell(
        row=year_total_row, column=6,
        value=f"=IF(B{year_total_row}=0,0,C{year_total_row}/B{year_total_row})",
    )
    all_avg_month = ws.cell(
        row=year_total_row, column=7,
        value=f"=AVERAGE(G{first_year_row}:G{last_year_row})",
    )
    all_months = ws.cell(row=year_total_row, column=8,
                         value=f"=SUM(H{first_year_row}:H{last_year_row})")
    for col in range(1, 9):
        c = ws.cell(row=year_total_row, column=col)
        c.fill = fill(GREEN)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        c.border = Border(left=THIN, right=THIN,
                          top=Side(style="thin", color=GREEN),
                          bottom=Side(style="double", color=GREEN))
        c.protection = LOCKED
    all_count.number_format = "0"
    all_count.alignment = Alignment(horizontal="center", vertical="center")
    all_total.number_format = CURRENCY_FMT
    all_total.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    all_share.number_format = "0.0%"
    all_share.alignment = Alignment(horizontal="center", vertical="center")
    all_avg_purchase.number_format = CURRENCY_FMT
    all_avg_purchase.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    all_avg_month.number_format = CURRENCY_FMT
    all_avg_month.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    all_months.number_format = "0"
    all_months.alignment = Alignment(horizontal="center", vertical="center")

    purchase_band_row = year_total_row + 2
    ws.row_dimensions[purchase_band_row].height = 26
    merge_after_styling(
        ws, f"A{purchase_band_row}:{last_col_letter}{purchase_band_row}",
        value="SPEND BY PURCHASER",
        font=HEADER_FONT,
        fill_color=GREEN,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(left=THIN, right=THIN, top=THIN, bottom=THIN),
    )

    header_row = purchase_band_row + 1
    matrix_headers = ["Purchaser"] + [str(y) for y in years] + ["Total (GH₵)", "% of all time"]
    for i, text in enumerate(matrix_headers, start=1):
        c = ws.cell(row=header_row, column=i, value=text)
        c.fill = fill(GOLD_LIGHT)
        c.font = Font(name=FONT, size=9, bold=True, color="6B4F2A")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="center", vertical="center")

    first_purch_row = header_row + 1
    last_purch_row = first_purch_row + PURCHASER_SLOTS - 1
    for i in range(PURCHASER_SLOTS):
        row = first_purch_row + i
        settings_row = LIST_FIRST_ROW + i
        ws.row_dimensions[row].height = 20
        band = "FFFFFF" if i % 2 == 0 else "FBF9F2"

        name = ws.cell(row=row, column=1,
                       value=f'=IF(Settings!A{settings_row}="","",Settings!A{settings_row})')
        name.fill = fill(band)
        name.font = BODY_FONT
        name.border = ALL_BORDERS
        name.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        name.protection = LOCKED

        for j, year in enumerate(years):
            col = 2 + j
            formula = (
                f'=IF($A{row}="","",IF(\'{year_summary_sheet(year)}\'!$C$'
                f'{SUMMARY_FIRST_PURCH_ROW + i}=0,"",'
                f"'{year_summary_sheet(year)}'!$C${SUMMARY_FIRST_PURCH_ROW + i}))"
            )
            c = ws.cell(row=row, column=col, value=formula)
            c.fill = fill(band)
            c.font = BODY_FONT
            c.border = ALL_BORDERS
            c.number_format = CURRENCY_FMT
            c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
            c.protection = LOCKED

        total = ws.cell(
            row=row, column=total_col,
            value=(f'=IF($A{row}="","",SUM(B{row}:'
                   f'{get_column_letter(last_year_col)}{row}))'),
        )
        pct = ws.cell(
            row=row, column=pct_col,
            value=f'=IF(OR($A{row}="",$C${year_total_row}=0),0,'
                  f'{get_column_letter(total_col)}{row}/$C${year_total_row})',
        )
        for c in (total, pct):
            c.fill = fill(band)
            c.font = BODY_FONT
            c.border = ALL_BORDERS
            c.protection = LOCKED
        total.number_format = CURRENCY_FMT
        total.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        pct.number_format = "0.0%"
        pct.alignment = Alignment(horizontal="center", vertical="center")

    listed_row = last_purch_row + 1
    unlisted_row = listed_row + 1
    ws.row_dimensions[listed_row].height = 24
    listed_label = ws.cell(row=listed_row, column=1, value="Listed purchasers")
    listed_label.font = BODY_BOLD
    listed_label.fill = fill(CREAM)
    listed_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    listed_label.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    listed_label.protection = LOCKED

    for col in range(2, last_year_col + 1):
        c = ws.cell(row=listed_row, column=col,
                    value=f"=SUM({get_column_letter(col)}{first_purch_row}:"
                          f"{get_column_letter(col)}{last_purch_row})")
        c.fill = fill(CREAM)
        c.font = BODY_BOLD
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        c.number_format = CURRENCY_FMT
        c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        c.protection = LOCKED
    ws.cell(row=listed_row, column=total_col,
            value=f"=SUM({get_column_letter(total_col)}{first_purch_row}:"
                  f"{get_column_letter(total_col)}{last_purch_row})")
    ws.cell(row=listed_row, column=pct_col,
            value=f"=IF($C${year_total_row}=0,0,"
                  f"{get_column_letter(total_col)}{listed_row}/$C${year_total_row})")
    for col in (total_col, pct_col):
        c = ws.cell(row=listed_row, column=col)
        c.fill = fill(CREAM)
        c.font = BODY_BOLD
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        c.protection = LOCKED
    ws.cell(row=listed_row, column=total_col).number_format = CURRENCY_FMT
    ws.cell(row=listed_row, column=total_col).alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.cell(row=listed_row, column=pct_col).number_format = "0.0%"
    ws.cell(row=listed_row, column=pct_col).alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[unlisted_row].height = 22
    unlisted_label = ws.cell(row=unlisted_row, column=1,
                             value="Unlisted / blank purchaser")
    unlisted_label.font = Font(name=FONT, size=9, italic=True, color=GREY)
    unlisted_label.fill = fill("FFFFFF")
    unlisted_label.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    unlisted_label.border = Border(left=THIN, right=THIN, bottom=THIN)
    unlisted_label.protection = LOCKED

    for j, year in enumerate(years):
        col = 2 + j
        c = ws.cell(row=unlisted_row, column=col,
                    value=f"='{year_summary_sheet(year)}'!$C${SUMMARY_UNLISTED_ROW}")
        c.fill = fill("FFFFFF")
        c.font = Font(name=FONT, size=9, italic=True, color=GREY)
        c.border = Border(left=THIN, right=THIN, bottom=THIN)
        c.number_format = CURRENCY_FMT
        c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        c.protection = LOCKED
    ws.cell(row=unlisted_row, column=total_col,
            value=(f"=SUM(B{unlisted_row}:"
                  f"{get_column_letter(last_year_col)}{unlisted_row})"))
    ws.cell(row=unlisted_row, column=pct_col,
            value=f"=IF($C${year_total_row}=0,0,"
                  f"{get_column_letter(total_col)}{unlisted_row}/$C${year_total_row})")
    for col in (total_col, pct_col):
        c = ws.cell(row=unlisted_row, column=col)
        c.fill = fill("FFFFFF")
        c.font = Font(name=FONT, size=9, italic=True, color=GREY)
        c.border = Border(left=THIN, right=THIN, bottom=THIN)
        c.protection = LOCKED
    ws.cell(row=unlisted_row, column=total_col).number_format = CURRENCY_FMT
    ws.cell(row=unlisted_row, column=total_col).alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.cell(row=unlisted_row, column=pct_col).number_format = "0.0%"
    ws.cell(row=unlisted_row, column=pct_col).alignment = Alignment(horizontal="center", vertical="center")

    ws.conditional_formatting.add(
        f"A{unlisted_row}:{last_col_letter}{unlisted_row}",
        FormulaRule(
            formula=[f"ABS(${get_column_letter(total_col)}${unlisted_row})>0.005"],
            fill=fill(RED_FILL),
            font=Font(name=FONT, size=9, bold=True, color=RED_TEXT),
        ),
    )

    note_row = unlisted_row + 2
    ws.row_dimensions[note_row].height = 32
    merge_after_styling(
        ws, f"A{note_row}:{last_col_letter}{note_row}",
        value=("Unlisted amounts are included in the all-years total. Add the purchaser's name "
               "in Settings to itemise them; the total is unaffected."),
        font=NOTE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1),
        border=Border(top=HAIR),
    )

    ws.print_area = f"A1:{last_col_letter}{note_row}"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"

    ws.protection.sheet = True


# ---------------------------------------------------------------- settings
def build_settings(ws, years):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREY
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 34

    ws.row_dimensions[1].height = 34
    merge_after_styling(
        ws, "A1:B1",
        value="SETTINGS & LISTS",
        font=TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=Side(style="medium", color=GOLD)),
    )
    ws.row_dimensions[2].height = 22
    merge_after_styling(
        ws, "A2:B2",
        value=(f"{min(years)}–{max(years)}  ·  Ghana Cedi (GH₵)  ·  "
               "Weeks Monday to Sunday"),
        font=SUBTITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[3].height = 12

    merge_after_styling(
        ws, "A4:B4",
        value="HOW TO USE",
        font=Font(name=FONT, size=9, bold=True, color="8A6D3B"),
        fill_color=CREAM,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )

    instructions = [
        ("1.  Add each purchaser's name in column A below — the Add Purchase sidebar, mobile "
         "Form and log dropdown all read this list."),
        ("2.  Each year has a Log for entry, a Summary, and 12 monthly summaries. Enter "
         "purchases in the Log; the monthly and yearly tabs are calculated reports."),
        ("3.  One purchase per row. Multiple purchases on the same day are expected — use "
         "Fuel Log → Add Purchase, the mobile Form, or the next blank log row. Each log holds "
         "5,000 purchases."),
        ("4.  Enter date, purchaser, amount and car number. Car number is required. To correct "
         "a purchase, edit its row; every change is written to the Audit Log."),
        ("5.  Weeks run Monday–Sunday. A week crossing month-end is split across two monthly "
         "summaries; both parts still count in the year total."),
    ]
    for i, text in enumerate(instructions):
        row = 5 + i
        ws.row_dimensions[row].height = (30, 40, 40, 40, 36)[i]
        merge_after_styling(
            ws, f"A{row}:B{row}",
            value=text,
            font=Font(name=FONT, size=9.5, color=INK),
            alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1),
            fill_color="FFFFFF",
            border=Border(bottom=HAIR),
        )

    ws.row_dimensions[11].height = 24
    merge_after_styling(
        ws, "A11:B11",
        value="USING THIS IN GOOGLE SHEETS",
        font=Font(name=FONT, size=9, bold=True, color="8A6D3B"),
        fill_color=CREAM,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[12].height = 80
    merge_after_styling(
        ws, "A12:B12",
        value=("Upload to Google Drive → Open with → Google Sheets → File → Save as Google "
               "Sheets. Then add Code.gs and Sidebar.html in Extensions → Apps Script and run "
               "setupAudit (see SETUP_GOOGLE_SHEETS.md). Formulas, dropdowns and frozen rows "
               "carry over; cell protection does not."),
        font=Font(name=FONT, size=9.5, color=INK),
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1),
        fill_color="F7FAF8",
        border=Border(top=HAIR, bottom=HAIR),
    )
    ws.row_dimensions[13].height = 70
    merge_after_styling(
        ws, "A13:B13",
        value=("The script powers Add Purchase, stamps Entered by/at in columns F–G, and writes "
               "every change to the Audit Log. Protect summaries and Settings; protect the Audit "
               "Log with a warning, not a hard restriction, so the sidebar can still write. Leave "
               "log entry columns editable for corrections."),
        font=NOTE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1),
        border=Border(bottom=THIN),
    )

    ws.row_dimensions[14].height = 26
    list_headers = ["Purchaser Name", "Role / Vehicle (optional)"]
    for i, text in enumerate(list_headers, start=1):
        c = ws.cell(row=14, column=i, value=text)
        c.fill = fill(GREEN)
        c.font = HEADER_FONT
        c.border = HEADER_BORDER
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for row in range(LIST_FIRST_ROW, LIST_LAST_ROW + 1):
        ws.row_dimensions[row].height = 20
        band = "FFFFFF" if (row - LIST_FIRST_ROW) % 2 == 0 else "FBF9F2"
        for col in range(1, 3):
            c = ws.cell(row=row, column=col)
            c.fill = fill(band)
            c.border = ALL_BORDERS
            c.font = BODY_FONT
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            c.protection = UNLOCKED

    ws.row_dimensions[65].height = 12
    ws.row_dimensions[66].height = 24
    merge_after_styling(
        ws, "A66:B66",
        value="MOBILE FORM",
        font=Font(name=FONT, size=9, bold=True, color="8A6D3B"),
        fill_color=CREAM,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[67].height = 24
    for col, value, font in (
        (1, "Mobile form link", Font(name=FONT, size=9.5, bold=True, color=INK)),
        (2, "Run Fuel Log → Create mobile form",
         Font(name=FONT, size=9.5, italic=True, color=GREY)),
    ):
        c = ws.cell(row=67, column=col, value=value)
        c.font = font
        c.fill = fill("FFFFFF")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.protection = LOCKED
    ws.row_dimensions[68].height = 22
    for col, value in (
        (1, "Owner account"),
        (2, "Set by administrator"),
    ):
        c = ws.cell(row=68, column=col, value=value)
        c.font = Font(name=FONT, size=9, color=GREY)
        c.fill = fill("FBF9F2")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.protection = LOCKED

    ws.row_dimensions[69].height = 12
    ws.row_dimensions[70].height = 24
    merge_after_styling(
        ws, "A70:B70",
        value="WEEKLY REPORT",
        font=Font(name=FONT, size=9, bold=True, color="8A6D3B"),
        fill_color=CREAM,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[71].height = 24
    for col, value, font, locked in (
        (1, "Recipients", Font(name=FONT, size=9.5, bold=True, color=INK), True),
        (2, "", Font(name=FONT, size=9.5, color=INK), False),
    ):
        c = ws.cell(row=71, column=col, value=value)
        c.font = font
        c.fill = fill("FFFFFF")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.protection = LOCKED if locked else UNLOCKED
    ws.row_dimensions[72].height = 22
    for col, value in ((1, "Last sent"), (2, "")):
        c = ws.cell(row=72, column=col, value=value)
        c.font = Font(name=FONT, size=9, color=GREY)
        c.fill = fill("FBF9F2")
        c.border = ALL_BORDERS
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.protection = LOCKED
    ws.row_dimensions[73].height = 30
    merge_after_styling(
        ws, "A73:B73",
        value="Comma-separated recipient emails. Manual sends can override this list.",
        font=Font(name=FONT, size=8, italic=True, color=GREY),
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1),
    )

    ws.print_area = "A1:B73"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"

    ws.protection.sheet = True


def update_settings_years(ws, years):
    ws["A2"] = (f"{min(years)}–{max(years)}  ·  Ghana Cedi (GH₵)  ·  "
                "Weeks Monday to Sunday")
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)


# ---------------------------------------------------------------- audit log
def build_audit_log(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "6B7280"
    widths = {"A": 19, "B": 28, "C": 14, "D": 7, "E": 14, "F": 24,
              "G": 16, "H": 20, "I": 14, "J": 24, "K": 24, "L": 11}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 34
    merge_after_styling(
        ws, "A1:L1",
        value="AUDIT LOG   ·   ENTRY AND EDIT HISTORY",
        font=TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=Side(style="medium", color=GOLD)),
    )
    ws.row_dimensions[2].height = 22
    merge_after_styling(
        ws, "A2:L2",
        value=("One row per change, written automatically by the Apps Script trigger "
               "after import to Google Sheets  ·  Timestamps use the spreadsheet time zone"),
        font=SUBTITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
        border=Border(bottom=THIN),
    )
    ws.row_dimensions[3].height = 20
    merge_after_styling(
        ws, "A3:L3",
        value=("Do not type in this sheet. See SETUP_GOOGLE_SHEETS.md for installation "
               "and protection steps."),
        font=Font(name=FONT, size=8.5, italic=True, color=GREY),
        alignment=Alignment(horizontal="left", vertical="center", indent=1),
    )

    headers = ["Timestamp", "User", "Sheet", "Row", "Purchase date", "Purchaser",
               "Amount (GH₵)", "Car Number", "Field", "Old value",
               "New value", "Action"]
    ws.row_dimensions[4].height = 30
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=text)
        c.fill = fill(GREEN)
        c.font = HEADER_FONT
        c.border = HEADER_BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:L1000"
    ws.print_title_rows = "1:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"

    ws.protection.sheet = True
    ws.protection.autoFilter = False



# ---------------------------------------------------------------- assembly
def build_year(wb, year):
    log = wb.create_sheet(log_sheet(year))
    build_log(log, year)

    month_cells = {}
    for month_no, (tab, full) in enumerate(MONTHS, start=1):
        ws = wb.create_sheet(month_sheet(year, tab))
        month_cells[tab] = build_month_summary(ws, year, month_no, full)
    summary = wb.create_sheet(year_summary_sheet(year))
    build_year_summary(summary, year, month_cells)


def build_master(years, path):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = f"Fuel Purchase Log {min(years)}–{max(years)}"
    wb.properties.subject = "Daily fuel purchases with weekly, monthly and yearly totals"
    wb.properties.creator = "Office Work"
    wb.properties.keywords = "fuel, purchases, Ghana, GHS"

    for year in years:
        build_year(wb, year)

    dashboard = wb.create_sheet("All Years")
    build_all_years(dashboard, years)
    audit = wb.create_sheet("Audit Log")
    build_audit_log(audit)
    settings = wb.create_sheet("Settings")
    build_settings(settings, years)

    reorder_sheets(wb, years)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


def add_years(path, new_years):
    wb = load_workbook(path)
    existing = detect_years(wb)
    additions = sorted({y for y in new_years if y not in existing})
    if not additions:
        return []

    total_sheets = len(wb.sheetnames) + 14 * len(additions)
    if total_sheets > 200:
        raise SystemExit(
            f"Adding {additions} would create {total_sheets} sheets; Google Sheets "
            "supports about 200. Start a new decade workbook instead."
        )

    for year in additions:
        build_year(wb, year)

    all_years = sorted(existing + additions)
    if "All Years" in wb.sheetnames:
        dashboard = wb["All Years"]
        clear_sheet(dashboard)
    else:
        dashboard = wb.create_sheet("All Years")
    build_all_years(dashboard, all_years)

    if "Audit Log" not in wb.sheetnames:
        audit = wb.create_sheet("Audit Log")
        build_audit_log(audit)

    if "Settings" in wb.sheetnames:
        update_settings_years(wb["Settings"], all_years)

    reorder_sheets(wb, all_years)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return additions


def main():
    parser = argparse.ArgumentParser(description="Build the multi-year fuel purchase master.")
    parser.add_argument("--years", nargs="+", type=int, default=YEARS,
                        help="Years to include (default: 2026-2030)")
    parser.add_argument("--out", default=str(OUT), help="Output path")
    parser.add_argument("--force", action="store_true",
                        help="Overwrite an existing workbook (destroys entered data)")
    args = parser.parse_args()

    out = Path(args.out)
    if out.exists() and not args.force:
        print(f"{out} already exists. Use --force to rebuild from scratch; "
              "use add_year.py to preserve data.")
        return 1

    years = sorted(set(args.years))
    build_master(years, out)
    print(f"Wrote {out} ({len(years)} years, {14 * len(years) + 3} sheets)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
